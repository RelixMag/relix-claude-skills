#!/usr/bin/env python3
"""
process_ticket_buyers.py
=========================
Turns raw ticket-buyer export files (CSV or XLSX) into Beehiiv-ready import
CSVs, following the Relix "Process Email Uploads" rules.

Two subcommands:

  map     Convert ONE raw source file into a Beehiiv-mapped CSV, applying
          header mapping, dedup, bounce/unsub filtering, and state/country/zip
          normalization.

  merge   Combine several already-mapped CSVs (same Beehiiv header schema,
          e.g. produced by `map` with different --artist values) into a single
          deduped-by-email output. Use this for multi-source Lists like Dayglo.

Examples
--------
Single venue file, CSV with headers:
    python3 process_ticket_buyers.py map 20260402_Bearsville_TicketBuyers_2026Q1_OG.csv \\
        --list Bearsville --signup-source "Ticket Buyer" \\
        --output 20260402_Bearsville_TicketBuyers_2026Q1_UPLOADED.csv

BBNY no-header XLSX (positional columns, col 5 = zip needing cleanup):
    python3 process_ticket_buyers.py map 20260708_BBNY_TicketBuyers_2026Q2_OG.xlsx \\
        --list BBNY --format xlsx-nohdr \\
        --output 20260708_BBNY_TicketBuyers_2026Q2_UPLOADED.csv

Combining multiple Dayglo sources with per-source Artist tags:
    python3 process_ticket_buyers.py map ruthie_fri.xlsx --list Dayglo --artist RuthieBandit \\
        --format xlsx-nohdr --output /tmp/ruthie_fri_mapped.csv
    python3 process_ticket_buyers.py map homestead.csv --list Dayglo --artist Homestead \\
        --output /tmp/homestead_mapped.csv
    python3 process_ticket_buyers.py merge /tmp/ruthie_fri_mapped.csv /tmp/homestead_mapped.csv \\
        --output 20260701_Dayglo_TicketBuyers_2026Q2_UPLOADED.csv

See references/field_values.md in this skill for the full field/value reference
and the BBNY zip-cleaning rules this script implements.
"""

import argparse
import csv
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    openpyxl = None


# ── Canonical Beehiiv output header order (matches established _UPLOADED.csv
# convention — email is 3rd, not 1st) ───────────────────────────────────────
BEEHIIV_HEADERS = [
    "First Name", "Last Name", "email", "Phone Number",
    "Street Address", "Address Line 2", "City Name", "State", "Zip Code",
    "Country Code", "List", "Signup Source", "Artist",
]

# ── Header alias table: normalized-key -> canonical Beehiiv field ──────────
# Keys are lowercased with all non-alphanumeric characters stripped, so
# "First Name", "first_name", "firstName" all normalize to "firstname".
def _norm_key(h):
    return re.sub(r"[^a-z0-9]", "", h.lower())


HEADER_ALIASES = {
    "email": "email",
    "firstname": "First Name",
    "lastname": "Last Name",
    # Some venue exports (e.g. CEG "Customer Contacts Report") give one combined
    # "Customer" column formatted as "LAST,FIRST" instead of separate columns.
    "customer": "_combined_name",
    "name": "_combined_name",
    "phone": "Phone Number",
    "phonenumber": "Phone Number",
    "address": "Street Address",
    "address1": "Street Address",
    "streetaddress": "Street Address",
    "address2": "Address Line 2",
    "addressline2": "Address Line 2",
    "aptsuite": "Address Line 2",
    "address3": "_address3",  # merged into Address Line 2 if that's empty
    "city": "City Name",
    "cityname": "City Name",
    "state": "State",
    "zip": "Zip Code",
    "zipcode": "Zip Code",
    "postalcode": "Zip Code",
    "country": "Country Code",
    "countrycode": "Country Code",
    "countrycodealpha2": "Country Code",
    # Filter-only columns (not copied to output, just used to drop rows)
    "emailbounced": "_bounced",
    "bounced": "_bounced",
    "emailunsubscribed": "_unsubscribed",
    "unsubscribed": "_unsubscribed",
    "optedout": "_unsubscribed",
    "status": "_status",  # generic status column, checked for bounce/unsub keywords
    # Klaviyo-style "Email Marketing Consent" export column (SUBSCRIBED /
    # NEVER_SUBSCRIBED / UNSUBSCRIBED) — reuse the generic _status keyword
    # check so UNSUBSCRIBED rows get dropped like any other unsub column.
    "emailmarketingconsent": "_status",
}

# ── Known valid field values (see references/field_values.md for details) ──
KNOWN_LISTS = {
    "BBNY", "Bearsville", "CapitolTheatre", "FANS.COM", "GMP-MIA", "GMP-NOLA",
    "Hubspot", "Jerry Dance Party", "RelixAllSubs", "RRPH", "GarciasCHI",
    "BBLV", "BBNash", "Dayglo",
}
KNOWN_SIGNUP_SOURCES = {
    "Contest", "Ticket Buyer", "Shop", "Meta", "Signup_Contests", "Relix.com",
}

# ── US state + Canadian province normalization ──────────────────────────────
STATE_NAME_MAP = {
    "Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR",
    "California": "CA", "Colorado": "CO", "Connecticut": "CT", "Delaware": "DE",
    "Florida": "FL", "Georgia": "GA", "Hawaii": "HI", "Idaho": "ID",
    "Illinois": "IL", "Indiana": "IN", "Iowa": "IA", "Kansas": "KS",
    "Kentucky": "KY", "Louisiana": "LA", "Maine": "ME", "Maryland": "MD",
    "Massachusetts": "MA", "Michigan": "MI", "Minnesota": "MN",
    "Mississippi": "MS", "Missouri": "MO", "Montana": "MT", "Nebraska": "NE",
    "Nevada": "NV", "New Hampshire": "NH", "New Jersey": "NJ",
    "New Mexico": "NM", "New York": "NY", "North Carolina": "NC",
    "North Dakota": "ND", "Ohio": "OH", "Oklahoma": "OK", "Oregon": "OR",
    "Pennsylvania": "PA", "Rhode Island": "RI", "South Carolina": "SC",
    "South Dakota": "SD", "Tennessee": "TN", "Texas": "TX", "Utah": "UT",
    "Vermont": "VT", "Virginia": "VA", "Washington": "WA",
    "West Virginia": "WV", "Wisconsin": "WI", "Wyoming": "WY",
    "District of Columbia": "DC", "Puerto Rico": "PR", "Virgin Islands": "VI",
    # Canadian provinces & territories
    "Alberta": "AB", "British Columbia": "BC", "Manitoba": "MB",
    "New Brunswick": "NB", "Newfoundland and Labrador": "NL",
    "Northwest Territories": "NT", "Nova Scotia": "NS", "Nunavut": "NU",
    "Ontario": "ON", "Prince Edward Island": "PE", "Quebec": "QC",
    "Saskatchewan": "SK", "Yukon": "YT",
    "UNKNOWN": "", "Unknown": "", "unknown": "",
}
VALID_STATE_CODES = {
    "AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "FL", "GA", "HI", "ID",
    "IL", "IN", "IA", "KS", "KY", "LA", "ME", "MD", "MA", "MI", "MN", "MS",
    "MO", "MT", "NE", "NV", "NH", "NJ", "NM", "NY", "NC", "ND", "OH", "OK",
    "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT", "VT", "VA", "WA", "WV",
    "WI", "WY", "DC", "PR", "VI", "GU", "AS", "MP", "AE", "AP", "AA",
    "AB", "BC", "MB", "NB", "NL", "NT", "NS", "NU", "ON", "PE", "QC", "SK", "YT",
}
CANADIAN_PROVINCE_CODES = {
    "AB", "BC", "MB", "NB", "NL", "NT", "NS", "NU", "ON", "PE", "QC", "SK", "YT",
}
US_ONLY_STATE_CODES = VALID_STATE_CODES - CANADIAN_PROVINCE_CODES

# ── Country normalization ───────────────────────────────────────────────────
COUNTRY_NAME_MAP = {
    "United States": "US", "United States Of America": "US",
    "United States of America": "US", "Canada": "CA",
    "United Kingdom": "GB", "Great Britain": "GB", "Germany": "DE",
    "Japan": "JP", "France": "FR", "Sweden": "SE", "Switzerland": "CH",
    "Italy": "IT", "Spain": "ES", "Australia": "AU", "Netherlands": "NL",
    "Mexico": "MX", "New Zealand": "NZ", "Finland": "FI", "Norway": "NO",
    "Greece": "GR", "Israel": "IL", "Brazil": "BR", "Austria": "AT",
    "Slovenia": "SI", "Croatia": "HR", "Portugal": "PT", "Denmark": "DK",
    "Ireland": "IE", "Uruguay": "UY", "Colombia": "CO", "Taiwan": "TW",
    "Hong Kong": "HK", "Virgin Islands,U.S": "VI", "Virgin Islands,British": "VG",
}
COUNTRY_ALPHA3_MAP = {
    "USA": "US", "CAN": "CA", "GBR": "GB", "DEU": "DE", "ITA": "IT",
    "AUS": "AU", "PRT": "PT", "FRA": "FR", "SWE": "SE", "CHE": "CH",
    "NLD": "NL", "AUT": "AT", "NOR": "NO", "IRL": "IE", "ESP": "ES",
    "FIN": "FI", "JPN": "JP", "GRC": "GR", "DNK": "DK", "HRV": "HR",
    "NZL": "NZ", "BRA": "BR", "COL": "CO", "MEX": "MX",
}
VALID_COUNTRY_CODES = {
    "US", "CA", "GB", "DE", "JP", "FR", "SE", "CH", "IT", "ES", "AU", "NL",
    "MX", "NZ", "FI", "NO", "GR", "IL", "BR", "AT", "SI", "HR", "PT", "DK",
    "IE", "UY", "CO", "TW", "HK", "VI", "VG",
}


def normalize_state(raw):
    val = (raw or "").strip()
    if not val:
        return val
    if val in STATE_NAME_MAP:
        return STATE_NAME_MAP[val]
    if val.upper() in VALID_STATE_CODES:
        return val.upper()
    return val  # international, left as-is


def normalize_country(raw):
    val = (raw or "").strip()
    if not val:
        return val
    if val in COUNTRY_NAME_MAP:
        return COUNTRY_NAME_MAP[val]
    if val.upper() in COUNTRY_ALPHA3_MAP:
        return COUNTRY_ALPHA3_MAP[val.upper()]
    if val.upper() in VALID_COUNTRY_CODES:
        return val.upper()
    return val  # unrecognized, left as-is


def clean_zip(raw):
    """General Zip Code cleaning rules (see references/field_values.md).
    Originally written for BBNY's positional XLSX col 5, but the rules
    themselves aren't BBNY-specific — zero-pad short numeric zips, truncate
    zip+4 to 5 digits, blank 6-9 digit numeric garbage, keep alphanumeric
    UK/Canadian postal codes as-is. Applied to every header-based format
    (csv-headers, xlsx-headers) in cmd_map, in addition to BBNY's xlsx-nohdr
    which has always called this at load time."""
    if raw is None:
        return ""
    # openpyxl reads numeric cells as int/float (e.g. 11222.0), not strings —
    # normalize whole-number floats to plain digit strings before cleaning.
    if isinstance(raw, float):
        raw = int(raw) if raw.is_integer() else raw
    val = str(raw).strip()
    if not val:
        return ""
    if any(c.isalpha() for c in val):
        # UK/Canadian postal code: keep, uppercase, normalize whitespace
        return " ".join(val.upper().split())
    if "-" in val:
        prefix = val.split("-")[0]
        if prefix.isdigit():
            return prefix.zfill(5)[:5]
        return val  # doesn't match a documented pattern — preserve as-is
    if val.isdigit():
        if len(val) <= 5:
            return val.zfill(5)
        return ""  # 6-9 digit numeric garbage -> blank
    # Doesn't match any documented pattern (e.g. a stray space like "021 31")
    # — preserve rather than silently discard real data.
    return val


# Backward-compatible alias — some docs/examples refer to this by its
# original BBNY-specific name.
clean_bbny_zip = clean_zip


def is_truthy(val):
    return str(val).strip().lower() in ("yes", "y", "true", "1")


def _open_text(filepath):
    """Open a source file for reading, tolerating non-UTF-8 exports (common in
    ticket-buyer files with special characters in addresses/names)."""
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            f = open(filepath, newline="", encoding=encoding)
            f.read()
            f.seek(0)
            return f
        except (UnicodeDecodeError, UnicodeError):
            continue
    return open(filepath, newline="", encoding="utf-8", errors="replace")


def _split_combined_name(raw):
    """Split a combined 'LAST,FIRST' name field (seen in CEG-style reports)
    into (First Name, Last Name). Falls back gracefully if there's no comma."""
    raw = (raw or "").strip()
    if not raw:
        return "", ""
    if "," in raw:
        last, _, first = raw.partition(",")
        return first.strip(), last.strip()
    # No comma — can't reliably split, put the whole thing in Last Name
    # rather than guess.
    return "", raw


def _rows_from_header_dicts(header_dicts):
    """Shared mapping/filter logic for any header-based source (CSV or XLSX-
    with-headers): map source columns to canonical Beehiiv fields, split
    combined name columns, and drop bounced/unsubscribed rows."""
    rows = []
    for row in header_dicts:
        out = {h: "" for h in BEEHIIV_HEADERS}
        bounced = unsubscribed = False
        address3 = ""
        for src_header, value in row.items():
            key = _norm_key(src_header)
            canon = HEADER_ALIASES.get(key)
            if canon is None:
                continue
            value = (value or "").strip()
            if canon == "_bounced":
                bounced = bounced or is_truthy(value)
            elif canon == "_unsubscribed":
                unsubscribed = unsubscribed or is_truthy(value)
            elif canon == "_status":
                low = value.lower()
                if "bounce" in low:
                    bounced = True
                if "unsub" in low or "opt out" in low or "optout" in low:
                    unsubscribed = True
            elif canon == "_address3":
                address3 = value
            elif canon == "_combined_name":
                first, last = _split_combined_name(value)
                out["First Name"], out["Last Name"] = first, last
            elif canon in out:
                out[canon] = value
        if address3 and not out["Address Line 2"]:
            out["Address Line 2"] = address3
        if bounced or unsubscribed:
            continue
        rows.append(out)
    return rows


def load_csv_with_headers(filepath):
    """CSV with a header row: map columns by (fuzzy) header name."""
    with _open_text(filepath) as f:
        reader = csv.DictReader(f)
        return _rows_from_header_dicts(reader)


def _xlsx_cell_to_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_xlsx_with_headers(filepath):
    """XLSX with a header row (e.g. a CEG 'Customer Contacts Report' export):
    first row is headers, map columns by (fuzzy) header name."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to read .xlsx files (pip install openpyxl)")
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    excel_rows = list(ws.iter_rows(values_only=True))
    if not excel_rows:
        return []
    headers = [_xlsx_cell_to_str(h) for h in excel_rows[0]]
    header_dicts = []
    for excel_row in excel_rows[1:]:
        row = {headers[i]: _xlsx_cell_to_str(v) for i, v in enumerate(excel_row) if i < len(headers)}
        header_dicts.append(row)
    return _rows_from_header_dicts(header_dicts)


def load_xlsx_nohdr(filepath):
    """BBNY-style positional XLSX: col2=Last, col3=First, col4=email, col5=zip."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to read .xlsx files (pip install openpyxl)")
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = []
    for excel_row in ws.iter_rows(values_only=True):
        if excel_row is None:
            continue
        last_name = excel_row[2] if len(excel_row) > 2 else None
        first_name = excel_row[3] if len(excel_row) > 3 else None
        email = excel_row[4] if len(excel_row) > 4 else None
        zip_raw = excel_row[5] if len(excel_row) > 5 else None
        rows.append({
            "email": str(email).strip() if email else "",
            "First Name": str(first_name).strip() if first_name else "",
            "Last Name": str(last_name).strip() if last_name else "",
            "Zip Code": clean_bbny_zip(zip_raw),
        })
    return rows


def detect_format(filepath):
    """Best-effort default. XLSX defaults to the no-header positional BBNY
    layout since that's the more common legacy case — pass
    --format xlsx-headers explicitly for header-based Excel exports like a
    CEG 'Customer Contacts Report'."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext in (".xlsx", ".xls"):
        return "xlsx-nohdr"
    return "csv-headers"


def cmd_map(args):
    fmt = args.format or detect_format(args.input)

    if fmt == "xlsx-nohdr":
        raw_rows = load_xlsx_nohdr(args.input)
    elif fmt == "xlsx-headers":
        raw_rows = load_xlsx_with_headers(args.input)
    elif fmt == "csv-headers":
        raw_rows = load_csv_with_headers(args.input)
    else:
        raise ValueError(f"Unknown --format: {fmt}")

    if args.list not in KNOWN_LISTS:
        print(f"WARNING: '{args.list}' is not a known List value. Confirm this is "
              f"intentional before publishing (see references/field_values.md) — "
              f"new List values need sign-off, unlike Artist tags.", file=sys.stderr)
    if args.signup_source not in KNOWN_SIGNUP_SOURCES:
        print(f"WARNING: '{args.signup_source}' is not a known Signup Source value.",
              file=sys.stderr)

    seen_emails = set()
    written = skipped_no_email = skipped_dupe = 0
    out_rows = []

    for row in raw_rows:
        email = (row.get("email") or "").strip().lower()
        if not email:
            skipped_no_email += 1
            continue
        if email in seen_emails:
            skipped_dupe += 1
            continue
        seen_emails.add(email)

        out = {h: row.get(h, "") for h in BEEHIIV_HEADERS}
        out["email"] = email
        out["State"] = normalize_state(out.get("State", ""))
        out["Country Code"] = normalize_country(out.get("Country Code", ""))
        # Standardize Zip Code (zero-pad, truncate zip+4, blank numeric
        # garbage) for every header-based format. BBNY's xlsx-nohdr already
        # cleaned its zip at load time (clean_zip, formerly clean_bbny_zip),
        # so re-running it here is a no-op for that format.
        out["Zip Code"] = clean_zip(out.get("Zip Code", ""))
        # If Country Code is still blank but State clearly identifies a US
        # jurisdiction (not Canadian), default Country Code to US — matches
        # established practice on files with no country column (or a
        # present-but-always-blank one) where every address was US-based.
        # Files with genuinely no state data (e.g. BBNY) are unaffected,
        # since State stays blank there and this never fires.
        if not out["Country Code"] and out["State"] in US_ONLY_STATE_CODES:
            out["Country Code"] = "US"
        # If Country Code and State are both still blank, fall back to the
        # Zip Code shape: a plain 5-digit numeric zip (after cleaning above)
        # is essentially always a US zip — Canadian postal codes are
        # alphanumeric (e.g. "M5V 3L9") and never match this pattern, so this
        # doesn't risk repeating the earlier Toronto/Ontario mislabeling
        # mistake. Skipped for BBNY (xlsx-nohdr), which has an explicit,
        # documented decision to leave Country Code blank regardless of zip.
        if not out["Country Code"] and fmt != "xlsx-nohdr" and re.fullmatch(r"\d{5}", out["Zip Code"] or ""):
            out["Country Code"] = "US"
        out["List"] = args.list
        out["Signup Source"] = args.signup_source
        out["Artist"] = args.artist or ""
        out_rows.append(out)
        written += 1

    with open(args.output, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=BEEHIIV_HEADERS)
        writer.writeheader()
        writer.writerows(out_rows)

    print(f"Output: {args.output}")
    print(f"  Format detected/used:  {fmt}")
    print(f"  Rows written:          {written}")
    print(f"  Skipped (no email):    {skipped_no_email}")
    print(f"  Skipped (duplicate):   {skipped_dupe}")
    print("  (Bounced/unsubscribed rows, if any columns were present, were already excluded.)")


def cmd_merge(args):
    seen_emails = set()
    written = skipped_dupe = 0
    out_rows = []

    for path in args.inputs:
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                email = (row.get("email") or "").strip().lower()
                if not email or email in seen_emails:
                    skipped_dupe += 1
                    continue
                seen_emails.add(email)
                out_rows.append({h: row.get(h, "") for h in BEEHIIV_HEADERS})
                written += 1

    with open(args.output, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=BEEHIIV_HEADERS)
        writer.writeheader()
        writer.writerows(out_rows)

    print(f"Output: {args.output}")
    print(f"  Source files merged:   {len(args.inputs)}")
    print(f"  Rows written:          {written}")
    print(f"  Skipped (duplicate/no email across sources): {skipped_dupe}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_map = sub.add_parser("map", help="Map one raw source file into a Beehiiv CSV")
    p_map.add_argument("input", help="Path to the raw source file (CSV or XLSX)")
    p_map.add_argument("--list", required=True, help="Beehiiv List value (e.g. BBNY, Bearsville)")
    p_map.add_argument("--signup-source", default="Ticket Buyer", dest="signup_source")
    p_map.add_argument("--artist", default="", help="Beehiiv Artist tag (freeform, optional)")
    p_map.add_argument("--format", choices=["xlsx-nohdr", "xlsx-headers", "csv-headers"], default=None,
                        help="Override auto-detection (auto-detects by extension otherwise)")
    p_map.add_argument("--output", required=True, help="Output CSV path")
    p_map.set_defaults(func=cmd_map)

    p_merge = sub.add_parser("merge", help="Merge multiple mapped CSVs, dedup by email")
    p_merge.add_argument("inputs", nargs="+", help="Paths to already-mapped CSVs")
    p_merge.add_argument("--output", required=True, help="Output CSV path")
    p_merge.set_defaults(func=cmd_merge)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
